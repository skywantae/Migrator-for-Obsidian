"""
엑셀(.xlsx) -> 옵시디언(Obsidian) 마이그레이션

- 시트 하나가 노트 하나가 됩니다.
- 이름이 비슷한 시트끼리 폴더로 묶습니다 (DAILY REPORT (...) 들, Visit WEEK NN 들).
- 표는 마크다운 표로, 시트에 박힌 그림은 attachments 폴더로 꺼냅니다.
- 통합문서마다 전체 시트를 훑는 색인 노트를 함께 만듭니다.

원본 파일은 읽기만 합니다. 열어서 읽을 뿐 고치거나 지우지 않습니다.
"""

import re
import time
from collections import Counter
from datetime import date, datetime, time as dtime
from pathlib import Path

from migrator_common import Stopped

INVALID_CHARS = re.compile(r'[\\/:*?"<>|\x00-\x1f]')

# 셀 안 줄바꿈은 마크다운 표에 그대로 못 넣는다
CELL_BREAK = "<br>"

# 칸 하나가 이보다 길면 마크다운 표에서 한 줄이 감당하기 어려워진다
WIDE_CELL = 200

# 이 정도로 긴 칸이 하나라도 있으면 표는 이미 못 읽는다
HUGE_CELL = 1000

# 긴 칸이 든 줄이 이 비율을 넘으면, 그 시트는 표가 아니라 글에 가깝다
WIDE_ROW_RATIO = 0.2

# 문단으로 펼 때, 이보다 짧은 값은 '이름: 값' 한 줄로 붙인다
INLINE_VALUE = 80

# 시트 이름 끝의 번호·날짜를 떼어 묶음 이름을 만든다
TRAILING_KEY = re.compile(r"[\s_\-.()\[\]#]*\d[\d\s\-_./]*[\s_\-.()\[\]#]*$")

IMAGE_EXT = {"png", "jpeg", "jpg", "gif", "bmp", "tiff", "emf", "wmf"}


# ============================== 값 다루기 ==============================
def sanitize_filename(name: str, fallback: str = "이름없음") -> str:
    name = INVALID_CHARS.sub(" ", name or "")
    name = re.sub(r"\s+", " ", name).strip().strip(".")
    return name[:110] if name else fallback


def yaml_escape(text: str) -> str:
    return str(text).replace("\\", "\\\\").replace('"', '\\"').replace("\n", " ")


def cell_text(value) -> str:
    """셀 값 하나를 사람이 읽을 글자로. 표 안에 넣어도 깨지지 않게 만든다."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        # 엑셀은 날짜를 datetime 으로 준다. 시각이 0시면 날짜만 쓴다.
        if (value.hour, value.minute, value.second) == (0, 0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, dtime):
        return value.strftime("%H:%M")
    if isinstance(value, bool):
        return "예" if value else "아니오"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    text = str(value).strip()
    text = text.replace("|", "\\|")          # 표 문법을 깨뜨리지 않게
    return re.sub(r"\r?\n", CELL_BREAK, text)


def first_date(grid: list) -> str:
    """시트에서 처음 나오는 날짜. Frontmatter 의 date 로 쓴다."""
    for row in grid:
        for value in row:
            if isinstance(value, datetime):
                return value.strftime("%Y-%m-%d")
            if isinstance(value, date):
                return value.strftime("%Y-%m-%d")
    return ""


# ============================== 시트 -> 격자 ==============================
def trim_grid(grid: list) -> list:
    """완전히 빈 행과 열을 걷어낸다. 병합 때문에 생긴 빈 칸이 대부분이다."""
    def filled(v):
        return v is not None and str(v).strip() != ""

    rows = [r for r in grid if any(filled(c) for c in r)]
    if not rows:
        return []

    width = max(len(r) for r in rows)
    rows = [list(r) + [None] * (width - len(r)) for r in rows]
    cols = [i for i in range(width) if any(filled(r[i]) for r in rows)]
    return [[r[i] for i in cols] for r in rows]


def read_grid(ws) -> list:
    """시트를 값 격자로 읽는다.

    병합된 칸은 openpyxl 이 첫 칸에만 값을 주므로, 나머지는 빈 칸으로 남는다.
    그 빈 칸들이 만든 빈 행·열을 걷어내면 사람이 보던 모양에 가까워진다.
    """
    return trim_grid([list(row) for row in ws.iter_rows(values_only=True)])


def looks_like_header(row: list) -> bool:
    """첫 줄을 표의 제목 줄로 볼 수 있는가.

    숫자나 날짜가 섞여 있으면 그건 이미 데이터다. 한 칸만 찬 줄은 표 제목이 아니라
    시트 제목인 경우가 많아 제외한다.
    """
    filled = [c for c in row if c is not None and str(c).strip() != ""]
    if len(filled) < 2:
        return False
    return all(isinstance(c, str) and len(c.strip()) <= 60 for c in filled)


# ============================== 격자 -> 마크다운 ==============================
def prefers_sections(grid: list) -> bool:
    """이 시트는 표보다 문단이 나은가.

    긴 칸이 하나 섞였다고 표를 버리지는 않는다. 대부분의 줄이 긴 글을 담고 있을 때,
    또는 한 칸이 표를 통째로 망가뜨릴 만큼 길 때만 문단으로 편다.
    """
    if not grid:
        return False

    wide_rows = 0
    for row in grid:
        lengths = [len(cell_text(c)) for c in row]
        if lengths and max(lengths) > HUGE_CELL:
            return True
        if lengths and max(lengths) > WIDE_CELL:
            wide_rows += 1
    return wide_rows > len(grid) * WIDE_ROW_RATIO


def grid_to_markdown(grid: list) -> str:
    if not grid:
        return "*(빈 시트)*"

    if prefers_sections(grid):
        return grid_to_sections(grid)

    rows = [[cell_text(c) for c in row] for row in grid]
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]

    if looks_like_header(grid[0]):
        header, body = rows[0], rows[1:]
    else:
        header, body = [""] * width, rows

    out = ["| " + " | ".join(header) + " |",
           "| " + " | ".join(["---"] * width) + " |"]
    out += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(out)


def grid_to_sections(grid: list) -> str:
    """긴 글이 든 시트는 표 대신 '항목 이름 + 내용' 문단으로 편다."""
    header = [cell_text(c) for c in grid[0]] if looks_like_header(grid[0]) else []
    body = grid[1:] if header else grid

    blocks = []
    for row in body:
        cells = [cell_text(c) for c in row]
        if not any(cells):
            continue
        lines = []
        for i, text in enumerate(cells):
            if not text:
                continue
            text = text.replace(CELL_BREAK, "\n")
            label = header[i] if i < len(header) and header[i] else ""
            if not label:
                lines.append(text)
            elif len(text) <= INLINE_VALUE and "\n" not in text:
                lines.append(f"**{label}**: {text}")      # 짧은 값은 한 줄로
            else:
                lines.append(f"**{label}**\n\n{text}")
        blocks.append("\n\n".join(lines))
    return "\n\n---\n\n".join(blocks) if blocks else "*(빈 시트)*"


# ============================== 시트 묶기 ==============================
def group_key(sheet_name: str) -> str:
    """시트 이름 끝의 번호·날짜를 떼어 묶음 이름을 만든다.

    'DAILY REPORT (20260824)' 와 'DAILY REPORT (20260820)' 은 같은 묶음이 되고,
    'CONTACT POINT' 처럼 뗄 것이 없으면 이름 그대로다.
    """
    name = re.sub(r"\s+", " ", (sheet_name or "").strip())
    stripped = TRAILING_KEY.sub("", name).strip(" _-.()[]#")
    return stripped or name


def plan_folders(sheet_names: list) -> dict:
    """{시트 이름: 들어갈 폴더 이름}. 혼자뿐인 묶음은 폴더를 만들지 않는다."""
    counts = Counter(group_key(n) for n in sheet_names)
    plan = {}
    for name in sheet_names:
        key = group_key(name)
        plan[name] = sanitize_filename(key, "") if counts[key] > 1 else ""
    return plan


def unique_note_name(base: str, used: set) -> str:
    """노트 이름은 볼트 전체에서 겹치면 안 된다 (겹치면 [[링크]]가 흐려진다)."""
    name, n = base, 2
    while name.lower() in used:
        name = f"{base} ({n})"
        n += 1
    used.add(name.lower())
    return name


# ============================== 그림 ==============================
def extract_images(ws, attachments_dir: Path, prefix: str) -> list:
    """시트에 박힌 그림을 파일로 꺼낸다. 꺼낸 파일 이름 목록을 돌려준다."""
    images = getattr(ws, "_images", None) or []
    if not images:
        return []

    attachments_dir.mkdir(parents=True, exist_ok=True)
    saved = []
    for i, img in enumerate(images, start=1):
        try:
            fmt = str(getattr(img, "format", "") or "png").lower()
            if fmt not in IMAGE_EXT:
                fmt = "png"
            name = sanitize_filename(f"{prefix}_{i}") + f".{fmt}"
            target = attachments_dir / name
            if not target.exists():
                data = img._data()
                if not data:
                    continue
                target.write_bytes(data)
            saved.append(name)
        except Exception:
            continue        # 그림 하나가 깨져도 나머지 시트는 계속 옮긴다
    return saved


# ============================== 노트 만들기 ==============================
def build_note(sheet_name: str, grid: list, images: list,
               source_name: str, index_note: str = "") -> str:
    lines = ["---", f'title: "{yaml_escape(sheet_name)}"']
    when = first_date(grid)
    if when:
        lines.append(f"date: {when}")
    lines.append(f'source_file: "{yaml_escape(source_name)}"')
    lines.append(f'sheet: "{yaml_escape(sheet_name)}"')
    lines.append("---")

    body = [grid_to_markdown(grid)]
    if images:
        body.append("\n".join(f"![[{n}]]" for n in images))
    if index_note:
        body.append(f"## 통합문서\n- [[{index_note}]]")

    return "\n".join(lines) + "\n\n" + "\n\n".join(body) + "\n"


def build_index(source_name: str, entries: list) -> str:
    """통합문서 색인 노트. 묶음별로 시트를 모아 보여준다."""
    by_folder = {}
    for note_name, folder in entries:
        by_folder.setdefault(folder or "", []).append(note_name)

    out = ["---", f'title: "{yaml_escape(source_name)}"',
           "type: workbook-index", "---", "",
           f"# {source_name}", "", f"시트 {len(entries)}개"]

    for folder in sorted(by_folder, key=lambda f: (f == "", f)):
        out.append("")
        out.append(f"## {folder}" if folder else "## 그 밖의 시트")
        out += [f"- [[{n}]]" for n in sorted(by_folder[folder])]
    return "\n".join(out) + "\n"


# ============================== 파일 살펴보기 ==============================
def inspect_workbook(path) -> dict:
    """옮기기 전에 무엇이 들어 있는지 알려준다 (GUI 의 [파일 확인])."""
    import openpyxl

    path = Path(path)
    if not path.exists():
        raise RuntimeError("파일을 찾을 수 없습니다.")

    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        names = list(wb.sheetnames)
        plan = plan_folders(names)
        images = sum(len(getattr(ws, "_images", None) or []) for ws in wb.worksheets)
        return {"name": path.name,
                "sheets": len(names),
                "folders": sorted({f for f in plan.values() if f}),
                "images": images}
    finally:
        wb.close()


# ============================== 전체 실행 ==============================
def run_excel_migration(settings: dict, log, progress, should_stop) -> dict:
    import openpyxl

    out_dir = Path(settings["out_dir"])
    out_dir.mkdir(parents=True, exist_ok=True)
    attachments_dir = out_dir / "attachments"
    files = [Path(p) for p in settings["files"]]
    want_images = settings.get("extract_images", True)
    skip_existing = settings.get("skip_existing", True)

    started = time.time()
    stats = {"saved": 0, "skipped": 0, "failed": 0, "images": 0}
    used_names = {p.stem.lower() for p in out_dir.rglob("*.md")}
    total_sheets = 0

    for path in files:
        log(f"\n{path.name} 을(를) 엽니다...")
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except Exception as e:
            log(f"  [에러] 열지 못했습니다: {e}")
            stats["failed"] += 1
            continue

        try:
            names = list(wb.sheetnames)
            plan = plan_folders(names)
            total_sheets += len(names)
            log(f"  시트 {len(names)}개 · 폴더 {len({f for f in plan.values() if f})}개")

            book_note = unique_note_name(sanitize_filename(path.stem), used_names)
            entries = []

            for i, sheet_name in enumerate(names, start=1):
                if should_stop():
                    raise Stopped()

                folder = plan[sheet_name]
                target_dir = out_dir / folder if folder else out_dir
                base = sanitize_filename(sheet_name, f"시트{i}")

                existing = target_dir / f"{base}.md"
                if skip_existing and existing.exists():
                    log(f"  [{i}/{len(names)}] 이미 있음 - 건너뜀: {sheet_name}")
                    entries.append((existing.stem, folder))
                    stats["skipped"] += 1
                    progress(i, len(names), sheet_name)
                    continue

                try:
                    ws = wb[sheet_name]
                    grid = read_grid(ws)
                    images = extract_images(ws, attachments_dir, base) if want_images else []
                    stats["images"] += len(images)

                    note_name = unique_note_name(base, used_names)
                    target_dir.mkdir(parents=True, exist_ok=True)
                    (target_dir / f"{note_name}.md").write_text(
                        build_note(sheet_name, grid, images, path.name, book_note),
                        encoding="utf-8")

                    entries.append((note_name, folder))
                    stats["saved"] += 1
                    log(f"  [{i}/{len(names)}] {sheet_name}"
                        + (f" · 그림 {len(images)}개" if images else ""))
                except Stopped:
                    raise
                except Exception as e:
                    log(f"  [{i}/{len(names)}] [에러] {sheet_name} -> {e}")
                    stats["failed"] += 1

                progress(i, len(names), sheet_name)

            if entries:
                (out_dir / f"{book_note}.md").write_text(
                    build_index(path.name, entries), encoding="utf-8")
                log(f"  색인 노트: {book_note}")
        finally:
            wb.close()

    log(f"\n시트 {total_sheets}개 중 {stats['saved']}개 저장"
        f" · 건너뜀 {stats['skipped']} · 실패 {stats['failed']}")
    return {"elapsed": time.time() - started, "saved": stats["saved"],
            "failed": stats["failed"], "skipped": stats["skipped"],
            "images": stats["images"]}
